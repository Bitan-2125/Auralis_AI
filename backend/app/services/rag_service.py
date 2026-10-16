"""Persistent basic RAG: document parsing, overlapping chunks and BM25 retrieval.

Runs locally using the application's database. No external vector service or
embedding download is required. Every query is restricted to one owned chat.
"""
import io
import math
import re
import zipfile
from collections import Counter
from pathlib import Path
from fastapi import HTTPException
from app import models
from app.config import settings

MAX_FILE_BYTES = settings.upload_max_mb * 1024 * 1024
MAX_TEXT_CHARS = 1000000
SUPPORTED = {".txt", ".md", ".csv", ".pdf", ".docx"}
STOPWORDS = set("a an the is are was were be to of and or in on for with this that it my me what how can please tell about from do does have document documents uploaded file files".split())


def extract_text(filename, raw):
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise HTTPException(415, "Use PDF, DOCX, TXT, Markdown or CSV files")
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise ValueError("Encrypted PDF")
            if len(reader.pages) > 500:
                raise HTTPException(413, "PDF exceeds the 500-page limit")
            parts, total = [], 0
            for page in reader.pages:
                text = page.extract_text() or ""
                total += len(text)
                if total > MAX_TEXT_CHARS:
                    raise HTTPException(413, "Extracted document is too large")
                parts.append(text)
            return "\n".join(parts)
        if suffix == ".docx":
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 30 * 1024 * 1024:
                    raise HTTPException(413, "Expanded DOCX is too large")
            from docx import Document
            doc = Document(io.BytesIO(raw))
            return "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
        return raw.decode("utf-8-sig")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "Could not read this file. Use a valid, unencrypted document or UTF-8 text. Scanned PDFs need OCR first.") from None


def chunk_text(text, chunk_size=1200, overlap=200):
    if len(text) > MAX_TEXT_CHARS:
        raise HTTPException(413, "Extracted document is too large")
    text = re.sub(r"\s+", " ", text).strip()
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        if end < len(text):
            boundary = text.rfind(" ", start + chunk_size // 2, end)
            if boundary > start:
                end = boundary
        chunks.append(text[start:end])
        if end == len(text):
            break
        start = max(start + 1, end - overlap)
    return chunks


def tokens(text):
    return [w for w in re.findall(r"\w+", text.casefold()) if w not in STOPWORDS and len(w) > 1]


def query(db, chat_id, question, n_results=4):
    rows = db.query(models.DocumentChunk, models.Document).join(models.Document).filter(models.Document.chat_id == chat_id).all()
    terms = set(tokens(question))
    if not rows:
        return []
    counts = [Counter(tokens(chunk.content + " " + doc.filename)) for chunk, doc in rows]
    lengths = [sum(c.values()) for c in counts]
    avg_length = max(1, sum(lengths) / len(rows))
    doc_freq = {term: sum(term in c for c in counts) for term in terms}
    ranked = []
    for idx, count in enumerate(counts):
        score = 0
        for term in terms:
            freq = count[term]
            if freq:
                idf = math.log(1 + (len(rows) - doc_freq[term] + 0.5) / (doc_freq[term] + 0.5))
                score += idf * (freq * 2.5) / (freq + 1.5 * (0.25 + 0.75 * lengths[idx] / avg_length))
        if score > 0:
            ranked.append((score, idx))
    ranked.sort(reverse=True)
    # Broad requests such as "summarize this document" have no lexical overlap.
    # Use opening passages across the attached files for these explicit requests.
    if not ranked and re.search(r"\b(summarize|summary|overview|summarise)\b", question, re.I):
        seen_documents = set()
        for idx, (chunk, document) in sorted(enumerate(rows), key=lambda item: item[1][0].chunk_index):
            if document.id not in seen_documents:
                ranked.append((0, idx))
                seen_documents.add(document.id)
    sources = []
    for _, idx in ranked[:n_results]:
        chunk, doc = rows[idx]
        sources.append({"number": len(sources) + 1, "document_id": doc.id,
                        "filename": doc.filename, "chunk_index": chunk.chunk_index,
                        "text": chunk.content, "excerpt": chunk.content[:240]})
    return sources
